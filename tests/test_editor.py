import tempfile
import zipfile
import os
from io import BytesIO
from openpyxl import Workbook, load_workbook
from xmlppt import editor as edmod
from xmlppt import PowerPointEditor
from lxml import etree


P_NS = edmod.P_NS
A_NS = edmod.A_NS
R_NS = edmod.R_NS


def _make_minimal_pptx(path: str) -> None:
    """Create a minimal pptx zip with one slide containing a textbox and a table."""
    # Content types
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        '</Types>'
    ).encode("utf-8")

    # presentation.xml
    presentation = (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:presentation xmlns:p="{P_NS}" xmlns:r="{R_NS}">'
        f'<p:sldIdLst>'
        f'<p:sldId id="256" r:id="rId1"/>'
        f'</p:sldIdLst>'
        f'</p:presentation>'
    ).encode("utf-8")

    presentation_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>'
        '</Relationships>'
    ).encode("utf-8")

    # slide1.xml with a textbox named TestBox and a table named MyTable
    slide = (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:sld xmlns:p="{P_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
        f'<p:cSld><p:spTree>'
        # Textbox
        f'<p:sp>'
        f'  <p:nvSpPr><p:cNvPr id="2" name="TestBox"/></p:nvSpPr>'
        f'  <p:txBody><a:bodyPr/><a:lstStyle/>'
        f'    <a:p><a:r><a:rPr/><a:t>Hello</a:t></a:r></a:p>'
        f'  </p:txBody>'
        f'</p:sp>'
        # Table in graphicFrame
        f'<p:graphicFrame>'
        f'  <p:nvGraphicFramePr><p:cNvPr id="5" name="MyTable"/></p:nvGraphicFramePr>'
        f'  <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
        f'    <a:tbl>'
        f'      <a:tr>'
        f'        <a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>R1C1</a:t></a:r></a:p></a:txBody></a:tc>'
        f'        <a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>R1C2</a:t></a:r></a:p></a:txBody></a:tc>'
        f'      </a:tr>'
        f'      <a:tr>'
        f'        <a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>R2C1</a:t></a:r></a:p></a:txBody></a:tc>'
        f'        <a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>R2C2</a:t></a:r></a:p></a:txBody></a:tc>'
        f'      </a:tr>'
        f'    </a:tbl>'
        f'  </a:graphicData></a:graphic>'
        f'</p:graphicFrame>'
        f'</p:spTree></p:cSld></p:sld>'
    ).encode("utf-8")

    # create zip
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('ppt/presentation.xml', presentation)
        z.writestr('ppt/_rels/presentation.xml.rels', presentation_rels)
        z.writestr('ppt/slides/slide1.xml', slide)


def _make_minimal_pptx_with_chart(path: str) -> None:
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        '<Override PartName="/ppt/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
        '<Override PartName="/ppt/charts/style1.xml" ContentType="application/vnd.ms-office.chartstyle+xml"/>'
        '<Override PartName="/ppt/charts/colors1.xml" ContentType="application/vnd.ms-office.chartcolorstyle+xml"/>'
        '<Override PartName="/ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/>'
        '</Types>'
    ).encode('utf-8')

    presentation = (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:presentation xmlns:p="{P_NS}" xmlns:r="{R_NS}">'
        f'<p:sldIdLst>'
        f'<p:sldId id="256" r:id="rId1"/>'
        f'</p:sldIdLst>'
        f'</p:presentation>'
    ).encode('utf-8')

    presentation_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>'
        '</Relationships>'
    ).encode('utf-8')

    slide = (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:sld xmlns:p="{P_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
        f'<p:cSld><p:spTree>'
        f'<p:graphicFrame>'
        f'  <p:nvGraphicFramePr><p:cNvPr id="5" name="TestChart"/></p:nvGraphicFramePr>'
        f'  <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        f'    <c:chart xmlns:c="{edmod.C_NS}" xmlns:r="{R_NS}" r:id="rId1"/>'
        f'  </a:graphicData></a:graphic>'
        f'</p:graphicFrame>'
        f'</p:spTree></p:cSld></p:sld>'
    ).encode('utf-8')

    chart_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<c:chartSpace xmlns:c="{edmod.C_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
        f'  <c:chart>'
        f'    <c:plotArea>'
        f'      <c:barChart>'
        f'        <c:ser>'
        f'          <c:cat>'
        f'            <c:strRef>'
        f'              <c:f>Sheet1!$A$2:$A$3</c:f>'
        f'              <c:strCache>'
        f'                <c:ptCount val="2"/>'
        f'                <c:pt idx="0"><c:v>One</c:v></c:pt>'
        f'                <c:pt idx="1"><c:v>Two</c:v></c:pt>'
        f'              </c:strCache>'
        f'            </c:strRef>'
        f'          </c:cat>'
        f'          <c:val>'
        f'            <c:numRef>'
        f'              <c:f>Sheet1!$B$2:$B$3</c:f>'
        f'              <c:numCache>'
        f'                <c:ptCount val="2"/>'
        f'                <c:pt idx="0"><c:v>10</c:v></c:pt>'
        f'                <c:pt idx="1"><c:v>20</c:v></c:pt>'
        f'              </c:numCache>'
        f'            </c:numRef>'
        f'          </c:val>'
        f'        </c:ser>'
        f'      </c:barChart>'
        f'    </c:plotArea>'
        f'  </c:chart>'
        f'</c:chartSpace>'
    ).encode('utf-8')

    chart_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="../embeddings/Microsoft_Excel_Worksheet1.xlsx"/>'
        '<Relationship Id="rId2" Type="http://schemas.microsoft.com/office/2011/relationships/chartStyle" Target="style1.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.microsoft.com/office/2011/relationships/chartColorStyle" Target="colors1.xml"/>'
        '</Relationships>'
    ).encode('utf-8')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws['A1'] = 'Category'
    ws['B1'] = 'Value'
    ws['A2'] = 'One'
    ws['B2'] = 10
    out = BytesIO()
    wb.save(out)
    workbook_bytes = out.getvalue()

    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('ppt/presentation.xml', presentation)
        z.writestr('ppt/_rels/presentation.xml.rels', presentation_rels)
        z.writestr('ppt/slides/slide1.xml', slide)
        z.writestr('ppt/slides/_rels/slide1.xml.rels', (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="../charts/chart1.xml"/>'
            '</Relationships>'
        ).encode('utf-8'))
        z.writestr('ppt/charts/chart1.xml', chart_xml)
        z.writestr('ppt/charts/_rels/chart1.xml.rels', chart_rels)
        z.writestr('ppt/charts/style1.xml', b'<cs:chartStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle"/>')
        z.writestr('ppt/charts/colors1.xml', b'<cs:colorStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle"/>')
        z.writestr('ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx', workbook_bytes)


def _make_minimal_pptx_with_sections(path: str) -> None:
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        '<Override PartName="/ppt/slides/slide2.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        '<Override PartName="/ppt/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
        '<Override PartName="/ppt/charts/style1.xml" ContentType="application/vnd.ms-office.chartstyle+xml"/>'
        '<Override PartName="/ppt/charts/colors1.xml" ContentType="application/vnd.ms-office.chartcolorstyle+xml"/>'
        '<Override PartName="/ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/>'
        '</Types>'
    ).encode('utf-8')

    presentation = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:presentation xmlns:p="{P_NS}" xmlns:r="{R_NS}" xmlns:p14="{edmod.P14_NS}">'
        '<p:sldIdLst>'
        '<p:sldId id="256" r:id="rId1"/>'
        '<p:sldId id="257" r:id="rId2"/>'
        '</p:sldIdLst>'
        '<p:extLst><p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
        '<p14:sectionLst>'
        '<p14:section name="main" id="{11111111-1111-1111-1111-111111111111}"><p14:sldIdLst><p14:sldId id="256"/></p14:sldIdLst></p14:section>'
        '<p14:section name="template_slides" id="{22222222-2222-2222-2222-222222222222}"><p14:sldIdLst><p14:sldId id="257"/></p14:sldIdLst></p14:section>'
        '</p14:sectionLst>'
        '</p:ext></p:extLst>'
        '</p:presentation>'
    ).encode('utf-8')

    presentation_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide2.xml"/>'
        '</Relationships>'
    ).encode('utf-8')

    slide1 = (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:sld xmlns:p="{P_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
        '<p:cSld><p:spTree><p:sp><p:nvSpPr><p:cNvPr id="2" name="Keep"/></p:nvSpPr></p:sp></p:spTree></p:cSld>'
        '</p:sld>'
    ).encode('utf-8')

    slide2 = (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:sld xmlns:p="{P_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
        '<p:cSld><p:spTree><p:graphicFrame>'
        '<p:nvGraphicFramePr><p:cNvPr id="5" name="TemplateChart"/></p:nvGraphicFramePr>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        f'<c:chart xmlns:c="{edmod.C_NS}" xmlns:r="{R_NS}" r:id="rId1"/>'
        '</a:graphicData></a:graphic>'
        '</p:graphicFrame></p:spTree></p:cSld>'
        '</p:sld>'
    ).encode('utf-8')

    slide2_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="../charts/chart1.xml"/>'
        '</Relationships>'
    ).encode('utf-8')

    chart_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<c:chartSpace xmlns:c="{edmod.C_NS}" xmlns:r="{R_NS}"><c:chart><c:plotArea/></c:chart></c:chartSpace>'
    ).encode('utf-8')

    chart_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="../embeddings/Microsoft_Excel_Worksheet1.xlsx"/>'
        '<Relationship Id="rId2" Type="http://schemas.microsoft.com/office/2011/relationships/chartStyle" Target="style1.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.microsoft.com/office/2011/relationships/chartColorStyle" Target="colors1.xml"/>'
        '</Relationships>'
    ).encode('utf-8')

    wb = Workbook()
    out = BytesIO()
    wb.save(out)

    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('ppt/presentation.xml', presentation)
        z.writestr('ppt/_rels/presentation.xml.rels', presentation_rels)
        z.writestr('ppt/slides/slide1.xml', slide1)
        z.writestr('ppt/slides/slide2.xml', slide2)
        z.writestr('ppt/slides/_rels/slide2.xml.rels', slide2_rels)
        z.writestr('ppt/charts/chart1.xml', chart_xml)
        z.writestr('ppt/charts/_rels/chart1.xml.rels', chart_rels)
        z.writestr('ppt/charts/style1.xml', b'<cs:chartStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle"/>')
        z.writestr('ppt/charts/colors1.xml', b'<cs:colorStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle"/>')
        z.writestr('ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx', out.getvalue())


def test_import():
    import xmlppt

    assert hasattr(xmlppt, 'PowerPointEditor')


def test_get_slide_returns_proxy(tmp_path):
    pptx = tmp_path / 'test_proxy.pptx'
    _make_minimal_pptx(str(pptx))

    editor = PowerPointEditor(str(pptx))
    slide = editor.get_slide(1)

    assert slide.slide_number == 1
    slide.edit_textbox('TestBox', 'Updated through proxy')
    assert 'Updated through proxy' in editor.files['ppt/slides/slide1.xml'].decode('utf-8')


def test_textbox_find_and_edit(tmp_path):
    pptx = tmp_path / 'test.pptx'
    _make_minimal_pptx(str(pptx))

    editor = PowerPointEditor(str(pptx))

    found = editor.find_textbox_anywhere('TestBox')
    assert found['slide_number'] == 1

    editor.edit_textbox_on_slide(1, 'TestBox', 'New\nLine')

    slide_xml = editor.files['ppt/slides/slide1.xml'].decode('utf-8')
    assert 'New' in slide_xml
    assert 'Line' in slide_xml


def test_drop_section_removes_section_slides_and_private_parts(tmp_path):
    pptx = tmp_path / 'test_sections.pptx'
    _make_minimal_pptx_with_sections(str(pptx))

    editor = PowerPointEditor(str(pptx))
    removed = editor.drop_section('template_slides')

    assert removed == [2]
    assert 'ppt/slides/slide1.xml' in editor.files
    assert 'ppt/slides/slide2.xml' not in editor.files
    assert 'ppt/slides/_rels/slide2.xml.rels' not in editor.files
    assert 'ppt/charts/chart1.xml' not in editor.files
    assert 'ppt/charts/style1.xml' not in editor.files
    assert 'ppt/charts/colors1.xml' not in editor.files
    assert 'ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx' not in editor.files

    presentation_xml = editor.files['ppt/presentation.xml'].decode('utf-8')
    presentation_rels = editor.files['ppt/_rels/presentation.xml.rels'].decode('utf-8')
    content_types = editor.files['[Content_Types].xml'].decode('utf-8')

    assert 'template_slides' not in presentation_xml
    assert 'slide2.xml' not in presentation_xml
    assert 'slide2.xml' not in presentation_rels
    assert '/ppt/slides/slide2.xml' not in content_types
    assert '/ppt/charts/chart1.xml' not in content_types


def test_table_find_and_edit_cell(tmp_path):
    pptx = tmp_path / 'test_table.pptx'
    _make_minimal_pptx(str(pptx))

    editor = PowerPointEditor(str(pptx))

    found = editor.find_table_anywhere('MyTable')
    assert found['slide_number'] == 1

    editor.edit_table_cell_on_slide(1, 'MyTable', 0, 1, 'NEWVAL')

    slide_xml = editor.files['ppt/slides/slide1.xml'].decode('utf-8')
    assert 'NEWVAL' in slide_xml


def test_table_edit_range(tmp_path):
    pptx = tmp_path / 'test_table2.pptx'
    _make_minimal_pptx(str(pptx))

    editor = PowerPointEditor(str(pptx))

    data = [['A1', 'A2'], ['B1', 'B2']]
    editor.edit_table_range_on_slide(1, 'MyTable', data)

    slide_xml = editor.files['ppt/slides/slide1.xml'].decode('utf-8')
    assert 'A1' in slide_xml and 'B2' in slide_xml


def test_chart_data_edit(tmp_path):
    pptx = tmp_path / 'test_chart.pptx'
    _make_minimal_pptx_with_chart(str(pptx))

    editor = PowerPointEditor(str(pptx))
    editor.edit_chart_data_on_slide(1, 'TestChart', ['NewOne', 'NewTwo'], [11, 22])

    chart_xml = editor.files['ppt/charts/chart1.xml'].decode('utf-8')
    assert 'NewOne' in chart_xml
    assert '22' in chart_xml
    assert 'Sheet1!$A$2:$A$3' in chart_xml
    assert 'Sheet1!$B$2:$B$3' in chart_xml

    rels = etree.fromstring(editor.files['ppt/charts/_rels/chart1.xml.rels'])
    workbook_target = rels.xpath(
        "./pr:Relationship[@Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/package']/@Target",
        namespaces=edmod.NS,
    )[0]
    workbook_path = PowerPointEditor._normalize_relationship_target('ppt/charts/chart1.xml', workbook_target)
    loaded_wb = load_workbook(BytesIO(editor.files[workbook_path]))
    sheet = loaded_wb['Sheet1']
    assert sheet['A2'].value == 'NewOne'
    assert sheet['B3'].value == 22


def test_embedded_workbook_for_chart_on_slide(tmp_path):
    pptx = tmp_path / 'test_embedded_chart.pptx'
    _make_minimal_pptx_with_chart(str(pptx))

    editor = PowerPointEditor(str(pptx))
    editor.edit_embedded_workbook_for_chart_on_slide(
        slide_number=1,
        chart_name='TestChart',
        categories=['CatA', 'CatB'],
        values=[100, 200],
        sheet_name='Sheet1',
    )

    chart_xml = editor.files['ppt/charts/chart1.xml'].decode('utf-8')
    assert 'CatA' in chart_xml
    assert '200' in chart_xml

    rels = etree.fromstring(editor.files['ppt/charts/_rels/chart1.xml.rels'])
    workbook_target = rels.xpath(
        "./pr:Relationship[@Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/package']/@Target",
        namespaces=edmod.NS,
    )[0]
    workbook_path = PowerPointEditor._normalize_relationship_target('ppt/charts/chart1.xml', workbook_target)
    workbook_data = editor.files[workbook_path]

    loaded_wb = load_workbook(BytesIO(workbook_data))
    sheet = loaded_wb['Sheet1']
    assert sheet['A2'].value == 'CatA'
    assert sheet['B3'].value == 200


def test_chart_data_multi_series_edit(tmp_path):
    pptx = tmp_path / 'test_multi_series_chart.pptx'
    _make_minimal_pptx_with_chart(str(pptx))

    editor = PowerPointEditor(str(pptx))
    chart_root = etree.fromstring(editor.files['ppt/charts/chart1.xml'])
    first_series = chart_root.xpath('.//c:barChart/c:ser', namespaces=edmod.NS)[0]
    first_series.getparent().append(etree.fromstring(etree.tostring(first_series)))
    editor.files['ppt/charts/chart1.xml'] = etree.tostring(chart_root, xml_declaration=True, encoding='UTF-8', standalone='yes')

    editor.edit_chart_data_multi_series_on_slide(
        slide_number=1,
        chart_name='TestChart',
        categories=['Q1', 'Q2', 'Q3'],
        series={
            'Revenue': [10, 20, 30],
            'Cost': [4, 8, 12],
        },
    )

    chart_xml = editor.files['ppt/charts/chart1.xml'].decode('utf-8')
    assert 'Sheet1!$A$2:$A$4' in chart_xml
    assert 'Sheet1!$B$2:$B$4' in chart_xml
    assert 'Sheet1!$C$2:$C$4' in chart_xml
    assert '30' in chart_xml
    assert '12' in chart_xml

    rels = etree.fromstring(editor.files['ppt/charts/_rels/chart1.xml.rels'])
    workbook_target = rels.xpath(
        "./pr:Relationship[@Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/package']/@Target",
        namespaces=edmod.NS,
    )[0]
    workbook_path = PowerPointEditor._normalize_relationship_target('ppt/charts/chart1.xml', workbook_target)
    loaded_wb = load_workbook(BytesIO(editor.files[workbook_path]))
    sheet = loaded_wb['Sheet1']

    assert sheet['A1'].value == 'Category'
    assert sheet['B1'].value == 'Revenue'
    assert sheet['C1'].value == 'Cost'
    assert sheet['A4'].value == 'Q3'
    assert sheet['B4'].value == 30
    assert sheet['C4'].value == 12


def test_duplicate_chart_slide_copies_chart_style_parts(tmp_path):
    pptx = tmp_path / 'test_duplicate_chart.pptx'
    _make_minimal_pptx_with_chart(str(pptx))

    editor = PowerPointEditor(str(pptx))
    new_slide_number = editor.duplicate_slide(1)

    assert new_slide_number == 2
    assert 'ppt/charts/chart2.xml' in editor.files
    assert 'ppt/charts/style2.xml' in editor.files
    assert 'ppt/charts/colors2.xml' in editor.files

    rels = etree.fromstring(editor.files['ppt/charts/_rels/chart2.xml.rels'])
    style_target = rels.xpath(
        "./pr:Relationship[@Type='http://schemas.microsoft.com/office/2011/relationships/chartStyle']/@Target",
        namespaces=edmod.NS,
    )
    color_target = rels.xpath(
        "./pr:Relationship[@Type='http://schemas.microsoft.com/office/2011/relationships/chartColorStyle']/@Target",
        namespaces=edmod.NS,
    )

    assert style_target == ['style2.xml']
    assert color_target == ['colors2.xml']

    content_types = etree.fromstring(editor.files['[Content_Types].xml'])
    assert content_types.xpath(
        "./ct:Override[@PartName='/ppt/charts/style2.xml']",
        namespaces=edmod.NS,
    )
    assert content_types.xpath(
        "./ct:Override[@PartName='/ppt/charts/colors2.xml']",
        namespaces=edmod.NS,
    )
