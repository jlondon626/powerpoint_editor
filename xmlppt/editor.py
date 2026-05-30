from zipfile import ZipFile
from lxml import etree
import re
import posixpath
from io import BytesIO
from openpyxl import load_workbook
import win32com.client
import os


# Namespaces
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CX_NS = "http://schemas.microsoft.com/office/drawing/2014/chartex"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"

NS = {
    "p": P_NS,
    "c": C_NS,
    "r": R_NS,
    "pr": PR_NS,
    "cx": CX_NS,
    "a": A_NS,
    "ct": CT_NS,
    "p14": P14_NS,
}


REL_TYPE_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
REL_TYPE_SLIDE_LAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
REL_TYPE_CHART = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
REL_TYPE_PACKAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
REL_TYPE_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
REL_TYPE_HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

SLIDE_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"

SLIDE_RE = re.compile(r"^ppt/slides/slide(\d+)\.xml$")


class PowerPointEditor:
    """A lightweight editor for .pptx files that exposes operations
    such as duplicating template slides, editing textboxes and charts,
    and saving the modified package back to disk.
    """

    def __init__(self, input_pptx: str):
        self.input_pptx = input_pptx
        self.files = self._load_pptx_files(input_pptx)
        self._has_changes = False

    # Basic package IO
    @staticmethod
    def _normalize_relationship_target(base_part: str, target: str) -> str:
        base_dir = posixpath.dirname(base_part)
        return posixpath.normpath(posixpath.join(base_dir, target))

    def _load_pptx_files(self, pptx_path: str) -> dict[str, bytes]:
        with ZipFile(pptx_path, "r") as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def _write_pptx_files(self, output_pptx: str) -> None:
        with ZipFile(output_pptx, "w") as archive:
            for name, data in self.files.items():
                archive.writestr(name, data)

    def _default_output_name(self, input_pptx: str) -> str:
        base, ext = os.path.splitext(input_pptx)
        return f"{base}_updated{ext}"

    def save(self, output_pptx: str | None = None) -> str:
        if output_pptx is None:
            output_pptx = self._default_output_name(self.input_pptx)

        self._write_pptx_files(output_pptx)
        self._has_changes = False
        return output_pptx

    # Package / part helpers
    def _rels_part_for(self, part_name: str) -> str:
        folder = posixpath.dirname(part_name)
        filename = posixpath.basename(part_name)
        return f"{folder}/_rels/{filename}.rels"

    def _next_numbered_part_name(self, folder: str, prefix: str, ext: str) -> str:
        pattern = re.compile(rf"^{re.escape(folder)}/{re.escape(prefix)}(\d+){re.escape(ext)}$")

        used = []
        for name in self.files:
            match = pattern.match(name)
            if match:
                used.append(int(match.group(1)))

        next_num = max(used, default=0) + 1
        return f"{folder}/{prefix}{next_num}{ext}"

    def _next_slide_part(self) -> tuple[int, str]:
        slide_numbers = [
            int(match.group(1))
            for name in self.files
            if (match := SLIDE_RE.match(name))
        ]

        next_num = max(slide_numbers, default=0) + 1
        return next_num, f"ppt/slides/slide{next_num}.xml"

    def _next_relationship_id(self, rels_root) -> str:
        used = []

        for rel in rels_root.xpath("./pr:Relationship", namespaces=NS):
            rid = rel.get("Id", "")
            match = re.match(r"rId(\d+)$", rid)
            if match:
                used.append(int(match.group(1)))

        return f"rId{max(used, default=0) + 1}"

    def _next_slide_id(self, presentation_root) -> str:
        ids = []

        for sld_id in presentation_root.xpath(".//p:sldId", namespaces=NS):
            value = sld_id.get("id")
            if value and value.isdigit():
                ids.append(int(value))

        return str(max(ids, default=255) + 1)

    def _add_content_type_override(self, part_name: str, content_type: str) -> None:
        content_types_part = "[Content_Types].xml"
        root = etree.fromstring(self.files[content_types_part])

        part_name_with_slash = f"/{part_name}"

        existing = root.xpath(
            f"./ct:Override[@PartName='{part_name_with_slash}']",
            namespaces=NS,
        )

        if existing:
            existing[0].set("ContentType", content_type)
        else:
            override = etree.SubElement(root, f"{{{CT_NS}}}Override")
            override.set("PartName", part_name_with_slash)
            override.set("ContentType", content_type)

        self.files[content_types_part] = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone="yes",
        )

    # Section-aware slide insertion
    def _find_insert_index_before_section(self, presentation_root, section_name: str) -> int | None:
        wanted = section_name.strip().casefold()

        sections = presentation_root.xpath(".//p14:section", namespaces=NS)
        if not sections:
            raise ValueError("No PowerPoint sections found in presentation.xml. Create a section named 'template_slides' in PowerPoint first.")

        target_section = None
        for section in sections:
            actual = (section.get("name") or "").strip().casefold()
            if actual == wanted:
                target_section = section
                break

        if target_section is None:
            raise ValueError(f"Section named '{section_name}' not found")

        section_slide_ids = target_section.xpath("./p14:sldIdLst/p14:sldId", namespaces=NS)
        if not section_slide_ids:
            raise ValueError(f"Section '{section_name}' contains no slides")

        first_template_slide_id = section_slide_ids[0].get("id")
        if not first_template_slide_id:
            raise ValueError(f"First slide in section '{section_name}' has no slide id")

        sld_id_lst = presentation_root.xpath("./p:sldIdLst", namespaces=NS)
        if not sld_id_lst:
            raise ValueError("No p:sldIdLst found in presentation.xml")

        all_slide_ids = sld_id_lst[0].xpath("./p:sldId", namespaces=NS)

        for index, sld_id in enumerate(all_slide_ids):
            if sld_id.get("id") == first_template_slide_id:
                return index

        raise ValueError(f"Could not locate first slide of section '{section_name}' in p:sldIdLst")

    def _add_slide_id_to_section_before(self, presentation_root, new_slide_id: str, before_section_name: str) -> None:
        wanted = before_section_name.strip().casefold()

        sections = presentation_root.xpath(".//p14:section", namespaces=NS)
        if not sections:
            return

        target_index = None
        for index, section in enumerate(sections):
            actual = (section.get("name") or "").strip().casefold()
            if actual == wanted:
                target_index = index
                break

        if target_index is None:
            raise ValueError(f"Section named '{before_section_name}' not found")

        if target_index == 0:
            raise ValueError("Cannot insert before section because it is the first section.")

        previous_section = sections[target_index - 1]

        sld_id_lst = previous_section.xpath("./p14:sldIdLst", namespaces=NS)
        if sld_id_lst:
            section_sld_id_lst = sld_id_lst[0]
        else:
            section_sld_id_lst = etree.SubElement(previous_section, f"{{{P14_NS}}}sldIdLst")

        new_section_sld_id = etree.SubElement(section_sld_id_lst, f"{{{P14_NS}}}sldId")
        new_section_sld_id.set("id", new_slide_id)

    def _add_slide_to_presentation(self, slide_part: str, before_section_name: str | None = None) -> int:
        presentation_part = "ppt/presentation.xml"
        presentation_rels_part = "ppt/_rels/presentation.xml.rels"

        presentation_root = etree.fromstring(self.files[presentation_part])
        presentation_rels_root = etree.fromstring(self.files[presentation_rels_part])

        new_rid = self._next_relationship_id(presentation_rels_root)

        rel = etree.SubElement(presentation_rels_root, f"{{{PR_NS}}}Relationship")
        rel.set("Id", new_rid)
        rel.set("Type", REL_TYPE_SLIDE)
        rel.set("Target", posixpath.relpath(slide_part, "ppt"))

        sld_id_lst = presentation_root.xpath("./p:sldIdLst", namespaces=NS)
        if not sld_id_lst:
            sld_id_lst_elem = etree.SubElement(presentation_root, f"{{{P_NS}}}sldIdLst")
        else:
            sld_id_lst_elem = sld_id_lst[0]

        new_slide_id = self._next_slide_id(presentation_root)

        new_sld_id = etree.Element(f"{{{P_NS}}}sldId")
        new_sld_id.set("id", new_slide_id)
        new_sld_id.set(f"{{{R_NS}}}id", new_rid)

        insert_index = None

        if before_section_name:
            insert_index = self._find_insert_index_before_section(presentation_root=presentation_root, section_name=before_section_name)

        if insert_index is None:
            sld_id_lst_elem.append(new_sld_id)
        else:
            sld_id_lst_elem.insert(insert_index, new_sld_id)

        if before_section_name:
            self._add_slide_id_to_section_before(presentation_root=presentation_root, new_slide_id=new_slide_id, before_section_name=before_section_name)

        self.files[presentation_part] = etree.tostring(presentation_root, xml_declaration=True, encoding="UTF-8", standalone="yes")
        self.files[presentation_rels_part] = etree.tostring(presentation_rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes")

        slide_num_match = SLIDE_RE.match(slide_part)
        if not slide_num_match:
            raise ValueError(f"Unexpected slide part name: {slide_part}")

        return int(slide_num_match.group(1))

    # Deep-copy slide dependencies
    def _copy_related_part_if_needed(self, source_part: str, dest_part: str, rel_type: str) -> str | None:
        if rel_type in {REL_TYPE_SLIDE_LAYOUT, REL_TYPE_IMAGE}:
            return None

        if rel_type == REL_TYPE_CHART:
            new_part = self._next_numbered_part_name("ppt/charts", "chart", ".xml")
        elif rel_type == REL_TYPE_PACKAGE:
            new_part = self._next_numbered_part_name("ppt/embeddings", "Microsoft_Excel_Worksheet", ".xlsx")
        else:
            return None

        self._deep_copy_part(source_part, new_part)

        dest_folder = posixpath.dirname(dest_part)
        return posixpath.relpath(new_part, dest_folder)

    def _deep_copy_part(self, source_part: str, dest_part: str) -> None:
        if source_part not in self.files:
            raise FileNotFoundError(f"Cannot copy missing part: {source_part}")

        self.files[dest_part] = self.files[source_part]

        source_rels_part = self._rels_part_for(source_part)
        dest_rels_part = self._rels_part_for(dest_part)

        if source_rels_part not in self.files:
            return

        rels_root = etree.fromstring(self.files[source_rels_part])

        for rel in rels_root.xpath("./pr:Relationship", namespaces=NS):
            target_mode = rel.get("TargetMode")
            if target_mode == "External":
                continue

            target = rel.get("Target")
            rel_type = rel.get("Type")

            if not target or not rel_type:
                continue

            source_related_part = self._normalize_relationship_target(source_part, target)

            if source_related_part not in self.files:
                continue

            new_target = self._copy_related_part_if_needed(source_part=source_related_part, dest_part=dest_part, rel_type=rel_type)

            if new_target:
                rel.set("Target", new_target)

        self.files[dest_rels_part] = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes")

    # Template slide lookup and duplication
    def find_slide_by_shape_name(self, shape_name: str) -> int:
        wanted = shape_name.strip().casefold()

        slide_parts = [(int(match.group(1)), name) for name in self.files if (match := SLIDE_RE.match(name))]
        slide_parts.sort(key=lambda item: item[0])

        for slide_number, slide_part in slide_parts:
            slide_root = etree.fromstring(self.files[slide_part])

            for shape in slide_root.xpath(".//p:sp", namespaces=NS):
                cNvPr = shape.xpath("./p:nvSpPr/p:cNvPr", namespaces=NS)
                if not cNvPr:
                    continue

                actual = (cNvPr[0].get("name") or "").strip().casefold()
                if actual == wanted:
                    return slide_number

            for frame in slide_root.xpath(".//p:graphicFrame", namespaces=NS):
                cNvPr = frame.xpath("./p:nvGraphicFramePr/p:cNvPr", namespaces=NS)
                if not cNvPr:
                    continue

                actual = (cNvPr[0].get("name") or "").strip().casefold()
                if actual == wanted:
                    return slide_number

        raise ValueError(f"Slide marker shape named '{shape_name}' not found")

    def remove_shape_on_slide(self, slide_number: int, shape_name: str) -> None:
        wanted = shape_name.strip().casefold()
        slide_part = f"ppt/slides/slide{slide_number}.xml"

        if slide_part not in self.files:
            raise FileNotFoundError(f"Slide not found: {slide_part}")

        slide_root = etree.fromstring(self.files[slide_part])
        removed = False

        for shape in slide_root.xpath(".//p:sp", namespaces=NS):
            cNvPr = shape.xpath("./p:nvSpPr/p:cNvPr", namespaces=NS)
            if not cNvPr:
                continue

            actual = (cNvPr[0].get("name") or "").strip().casefold()
            if actual == wanted:
                shape.getparent().remove(shape)
                removed = True
                break

        if not removed:
            for frame in slide_root.xpath(".//p:graphicFrame", namespaces=NS):
                cNvPr = frame.xpath("./p:nvGraphicFramePr/p:cNvPr", namespaces=NS)
                if not cNvPr:
                    continue

                actual = (cNvPr[0].get("name") or "").strip().casefold()
                if actual == wanted:
                    frame.getparent().remove(frame)
                    removed = True
                    break

        if removed:
            self.files[slide_part] = etree.tostring(slide_root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            self._has_changes = True

    def duplicate_slide(self, template_slide_number: int, before_section_name: str | None = None) -> int:
        source_slide_part = f"ppt/slides/slide{template_slide_number}.xml"

        if source_slide_part not in self.files:
            raise FileNotFoundError(f"Template slide not found: {source_slide_part}")

        new_slide_number, new_slide_part = self._next_slide_part()

        self._deep_copy_part(source_slide_part, new_slide_part)
        self._add_content_type_override(new_slide_part, SLIDE_CONTENT_TYPE)

        inserted_slide_number = self._add_slide_to_presentation(slide_part=new_slide_part, before_section_name=before_section_name)

        self._has_changes = True
        return inserted_slide_number

    def duplicate_template_slide(self, template_name: str, before_section_name: str = "template_slides") -> int:
        marker_shape_name = f"TEMPLATE__{template_name}"

        template_slide_number = self.find_slide_by_shape_name(marker_shape_name)

        new_slide_number = self.duplicate_slide(template_slide_number=template_slide_number, before_section_name=before_section_name)

        # Remove marker from generated slide so future lookups only find template slides.
        self.remove_shape_on_slide(slide_number=new_slide_number, shape_name=marker_shape_name)

        return new_slide_number

    # Textbox editing
    def _find_textbox_shape_in_slide(self, slide_root, textbox_name: str):
        wanted = textbox_name.strip().casefold()

        for shape in slide_root.xpath(".//p:sp", namespaces=NS):
            cNvPr = shape.xpath("./p:nvSpPr/p:cNvPr", namespaces=NS)
            if not cNvPr:
                continue

            actual_name = (cNvPr[0].get("name") or "").strip().casefold()
            if actual_name == wanted:
                return shape

        return None

    def _append_text_run(self, paragraph, text: str, bold: bool = False) -> None:
        run = etree.SubElement(paragraph, f"{{{A_NS}}}r")

        if bold:
            etree.SubElement(run, f"{{{A_NS}}}rPr", b="1")

        text_element = etree.SubElement(run, f"{{{A_NS}}}t")

        if text.startswith(" ") or text.endswith(" "):
            text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

        text_element.text = text

    def _replace_textbox_runs(self, txBody_elem, paragraphs: list[list[tuple[str, bool]]]) -> None:
        for paragraph in txBody_elem.xpath("./a:p", namespaces=NS):
            txBody_elem.remove(paragraph)

        for runs in paragraphs:
            paragraph = etree.SubElement(txBody_elem, f"{{{A_NS}}}p")

            if not runs:
                self._append_text_run(paragraph, "", False)
            else:
                for text, bold in runs:
                    self._append_text_run(paragraph, text, bold)

    def _text_to_paragraph_runs(self, text: str) -> list[list[tuple[str, bool]]]:
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        return [[(line, False)] for line in normalized.split("\n")]

    def _parse_textbox_markup(self, markup: str) -> list[list[tuple[str, bool]]]:
        normalized = markup.replace("\r\n", "\n").replace("\r", "\n")
        normalized = re.sub(r"(?i)<br\s*/?>", "\n", normalized)

        paragraphs: list[list[tuple[str, bool]]] = []

        for line in normalized.split("\n"):
            runs: list[tuple[str, bool]] = []
            last_index = 0

            for match in re.finditer(r"(?i)<b>(.*?)</b>", line, flags=re.DOTALL):
                if match.start() > last_index:
                    runs.append((line[last_index:match.start()], False))

                runs.append((match.group(1), True))
                last_index = match.end()

            if last_index < len(line):
                runs.append((line[last_index:], False))

            if not runs:
                runs.append(("", False))

            paragraphs.append(runs)

        return paragraphs

    def find_textbox_anywhere(self, textbox_name: str) -> dict:
        slide_parts = [(int(match.group(1)), name) for name in self.files if (match := SLIDE_RE.match(name))]
        slide_parts.sort(key=lambda item: item[0])

        for slide_number, slide_part in slide_parts:
            slide_root = etree.fromstring(self.files[slide_part])
            shape = self._find_textbox_shape_in_slide(slide_root, textbox_name)

            if shape is not None:
                return {
                    "slide_number": slide_number,
                    "slide_part": slide_part,
                }

        raise ValueError(f"Textbox named '{textbox_name}' not found anywhere in presentation")

    def edit_textbox_on_slide(self, slide_number: int, textbox_name: str, new_text: str) -> None:
        slide_part = f"ppt/slides/slide{slide_number}.xml"

        if slide_part not in self.files:
            raise FileNotFoundError(f"Slide not found: {slide_part}")

        slide_root = etree.fromstring(self.files[slide_part])
        shape = self._find_textbox_shape_in_slide(slide_root, textbox_name)

        if shape is None:
            raise ValueError(f"Textbox named '{textbox_name}' not found on slide {slide_number}")

        txBody = shape.xpath("./p:txBody", namespaces=NS)
        if not txBody:
            raise ValueError(f"No text body found in textbox '{textbox_name}'")

        paragraphs = self._text_to_paragraph_runs(new_text)
        self._replace_textbox_runs(txBody[0], paragraphs)

        self.files[slide_part] = etree.tostring(slide_root, xml_declaration=True, encoding="UTF-8", standalone="yes")

        self._has_changes = True

    def edit_textbox_html_on_slide(self, slide_number: int, textbox_name: str, html: str) -> None:
        slide_part = f"ppt/slides/slide{slide_number}.xml"

        if slide_part not in self.files:
            raise FileNotFoundError(f"Slide not found: {slide_part}")

        slide_root = etree.fromstring(self.files[slide_part])
        shape = self._find_textbox_shape_in_slide(slide_root, textbox_name)

        if shape is None:
            raise ValueError(f"Textbox named '{textbox_name}' not found on slide {slide_number}")

        txBody = shape.xpath("./p:txBody", namespaces=NS)
        if not txBody:
            raise ValueError(f"No text body found in textbox '{textbox_name}'")

        paragraphs = self._parse_textbox_markup(html)
        self._replace_textbox_runs(txBody[0], paragraphs)

        self.files[slide_part] = etree.tostring(slide_root, xml_declaration=True, encoding="UTF-8", standalone="yes")

        self._has_changes = True

    def edit_textbox_runs_on_slide(self, slide_number: int, textbox_name: str, paragraphs: list[list[tuple[str, bool]]]) -> None:
        slide_part = f"ppt/slides/slide{slide_number}.xml"

        if slide_part not in self.files:
            raise FileNotFoundError(f"Slide not found: {slide_part}")

        slide_root = etree.fromstring(self.files[slide_part])
        shape = self._find_textbox_shape_in_slide(slide_root, textbox_name)

        if shape is None:
            raise ValueError(f"Textbox named '{textbox_name}' not found on slide {slide_number}")

        txBody = shape.xpath("./p:txBody", namespaces=NS)
        if not txBody:
            raise ValueError(f"No text body found in textbox '{textbox_name}'")

        self._replace_textbox_runs(txBody[0], paragraphs)

        self.files[slide_part] = etree.tostring(slide_root, xml_declaration=True, encoding="UTF-8", standalone="yes")

        self._has_changes = True

    # Backward-compatible anywhere methods
    def edit_textbox(self, textbox_name: str, new_text: str) -> None:
        found = self.find_textbox_anywhere(textbox_name)
        self.edit_textbox_on_slide(found["slide_number"], textbox_name, new_text)

    def edit_textbox_html(self, textbox_name: str, html: str) -> None:
        found = self.find_textbox_anywhere(textbox_name)
        self.edit_textbox_html_on_slide(found["slide_number"], textbox_name, html)

    # Chart lookup/editing
    def find_chart_on_slide(self, slide_number: int, chart_name: str) -> dict:
        wanted = chart_name.strip().casefold()

        slide_part = f"ppt/slides/slide{slide_number}.xml"
        rels_part = f"ppt/slides/_rels/slide{slide_number}.xml.rels"

        if slide_part not in self.files:
            raise FileNotFoundError(f"Slide not found: {slide_part}")

        if rels_part not in self.files:
            raise FileNotFoundError(f"Slide relationships not found: {rels_part}")

        slide_root = etree.fromstring(self.files[slide_part])
        rels_root = etree.fromstring(self.files[rels_part])

        for frame in slide_root.xpath(".//p:graphicFrame", namespaces=NS):
            cNvPr = frame.xpath("./p:nvGraphicFramePr/p:cNvPr", namespaces=NS)
            if not cNvPr:
                continue

            actual_name = (cNvPr[0].get("name") or "").strip().casefold()
            if actual_name != wanted:
                continue

            chart_refs = frame.xpath(".//c:chart", namespaces=NS)
            chart_kind = "c"

            if not chart_refs:
                chart_refs = frame.xpath(".//cx:chart", namespaces=NS)
                chart_kind = "cx"

            if not chart_refs:
                continue

            rel_id = chart_refs[0].get(f"{{{R_NS}}}id")
            if not rel_id:
                continue

            rel = rels_root.xpath(f"./pr:Relationship[@Id='{rel_id}']", namespaces=NS)
            if not rel:
                continue

            target = rel[0].get("Target")
            if not target:
                continue

            return {
                "slide_number": slide_number,
                "slide_part": slide_part,
                "chart_part": self._normalize_relationship_target(slide_part, target),
                "rel_id": rel_id,
                "chart_kind": chart_kind,
            }

        raise ValueError(f"Chart named '{chart_name}' not found on slide {slide_number}")

    def find_chart_anywhere(self, chart_name: str) -> dict:
        slide_parts = [(int(match.group(1)), name) for name in self.files if (match := SLIDE_RE.match(name))]
        slide_parts.sort(key=lambda item: item[0])

        for slide_number, _slide_part in slide_parts:
            try:
                return self.find_chart_on_slide(slide_number, chart_name)
            except ValueError:
                continue

        raise ValueError(f"Chart named '{chart_name}' not found anywhere in presentation")

    def _replace_str_cache(self, str_cache, labels: list[str]) -> None:
        for pt in str_cache.xpath("./c:pt", namespaces=NS):
            str_cache.remove(pt)

        pt_count = str_cache.find(f"{{{C_NS}}}ptCount")
        if pt_count is None:
            pt_count = etree.Element(f"{{{C_NS}}}ptCount")
            str_cache.insert(0, pt_count)

        pt_count.set("val", str(len(labels)))

        for index, label in enumerate(labels):
            pt = etree.SubElement(str_cache, f"{{{C_NS}}}pt", idx=str(index))
            etree.SubElement(pt, f"{{{C_NS}}}v").text = str(label)

    def _replace_num_cache(self, num_cache, values: list[float]) -> None:
        for pt in num_cache.xpath("./c:pt", namespaces=NS):
            num_cache.remove(pt)

        pt_count = num_cache.find(f"{{{C_NS}}}ptCount")
        if pt_count is None:
            pt_count = etree.Element(f"{{{C_NS}}}ptCount")
            num_cache.insert(0, pt_count)

        pt_count.set("val", str(len(values)))

        for index, value in enumerate(values):
            pt = etree.SubElement(num_cache, f"{{{C_NS}}}pt", idx=str(index))
            etree.SubElement(pt, f"{{{C_NS}}}v").text = str(value)

    def _get_embedded_workbook_target(self, chart_part: str, rels_root) -> str | None:
        for rel in rels_root.xpath("./pr:Relationship", namespaces=NS):
            if rel.get("Type") == REL_TYPE_PACKAGE:
                return self._normalize_relationship_target(chart_part, rel.get("Target", ""))

        return None

    def _update_chartex_chart(self, chart_root, categories: list[str], values: list[float], subtotal_indices: list[int] | None = None) -> None:
        cat_lvl = chart_root.xpath("//cx:strDim[@type='cat']/cx:lvl", namespaces=NS)
        if cat_lvl:
            lvl = cat_lvl[0]
            lvl.set("ptCount", str(len(categories)))

            for pt in lvl.xpath("./cx:pt", namespaces=NS):
                lvl.remove(pt)

            for index, category in enumerate(categories):
                pt = etree.SubElement(lvl, f"{{{CX_NS}}}pt", {"idx": str(index)})
                pt.text = str(category)

            f_elem = lvl.getparent().xpath("./cx:f", namespaces=NS)
            if f_elem:
                f_elem[0].text = f"Sheet1!$A$2:$A${len(categories) + 1}"

        val_lvl = chart_root.xpath("//cx:numDim[@type='val']/cx:lvl", namespaces=NS)
        if val_lvl:
            lvl = val_lvl[0]
            lvl.set("ptCount", str(len(values)))

            for pt in lvl.xpath("./cx:pt", namespaces=NS):
                lvl.remove(pt)

            for index, value in enumerate(values):
                pt = etree.SubElement(lvl, f"{{{CX_NS}}}pt", {"idx": str(index)})
                pt.text = str(value)

            f_elem = lvl.getparent().xpath("./cx:f", namespaces=NS)
            if f_elem:
                f_elem[0].text = f"Sheet1!$B$2:$B${len(values) + 1}"

        if subtotal_indices is not None:
            subtotal_node = chart_root.xpath("//cx:layoutPr/cx:subtotals", namespaces=NS)
            if subtotal_node:
                subtotals = subtotal_node[0]

                for idx_elem in subtotals.xpath("./cx:idx", namespaces=NS):
                    subtotals.remove(idx_elem)

                for idx in subtotal_indices:
                    etree.SubElement(subtotals, f"{{{CX_NS}}}idx", {"val": str(idx)})

    def _update_regular_chart_cache(self, chart_root, categories: list[str], values: list[float]) -> None:
        str_caches = chart_root.xpath("//c:cat//c:strCache", namespaces=NS)
        if str_caches:
            self._replace_str_cache(str_caches[0], categories)

        num_caches = chart_root.xpath("//c:val//c:numCache", namespaces=NS)
        if num_caches:
            self._replace_num_cache(num_caches[0], values)

    def edit_waterfall_data_on_slide(self, slide_number: int, chart_name: str, categories: list[str], values: list[float]) -> None:
        if len(categories) != len(values):
            raise ValueError("categories and values must have the same length")

        found = self.find_chart_on_slide(slide_number, chart_name)
        chart_part = found["chart_part"]

        chart_root = etree.fromstring(self.files[chart_part])

        str_cache_nodes = chart_root.xpath(".//c:cat//c:strCache", namespaces=NS)
        num_cache_nodes = chart_root.xpath(".//c:val//c:numCache", namespaces=NS)

        if not str_cache_nodes:
            raise ValueError("No category string cache found")

        if not num_cache_nodes:
            raise ValueError("No value numeric cache found")

        self._replace_str_cache(str_cache_nodes[0], categories)
        self._replace_num_cache(num_cache_nodes[0], values)

        self.files[chart_part] = etree.tostring(chart_root, xml_declaration=True, encoding="UTF-8", standalone="yes")

        self._has_changes = True

    def edit_embedded_workbook_for_chart_on_slide(self, slide_number: int, chart_name: str, categories: list[str], values: list[float], sheet_name: str | None = None, subtotal_indices: list[int] | None = None) -> None:
        if len(categories) != len(values):
            raise ValueError("categories and values must have the same length")

        found = self.find_chart_on_slide(slide_number, chart_name)
        chart_part = found["chart_part"]

        rels_part = f"{posixpath.dirname(chart_part)}/_rels/{posixpath.basename(chart_part)}.rels"

        if rels_part not in self.files:
            raise FileNotFoundError(f"Chart rels part not found: {rels_part}")

        rels_root = etree.fromstring(self.files[rels_part])
        workbook_target = self._get_embedded_workbook_target(chart_part, rels_root)

        if not workbook_target:
            raise ValueError("No embedded workbook relationship found for this chart")

        if workbook_target not in self.files:
            raise FileNotFoundError(f"Embedded workbook not found: {workbook_target}")

        wb_bytes = BytesIO(self.files[workbook_target])
        wb = load_workbook(wb_bytes)
        sheet = wb[sheet_name] if sheet_name else wb.active

        start_row = 2
        end_row = max(sheet.max_row, len(categories) + 1)

        for row in range(start_row, end_row + 1):
            sheet[f"A{row}"] = None
            sheet[f"B{row}"] = None

        for row_index, (category, value) in enumerate(zip(categories, values), start=start_row):
            sheet[f"A{row_index}"] = category
            sheet[f"B{row_index}"] = value

        out_wb = BytesIO()
        wb.save(out_wb)
        self.files[workbook_target] = out_wb.getvalue()

        chart_root = etree.fromstring(self.files[chart_part])

        self._update_chartex_chart(chart_root=chart_root, categories=categories, values=values, subtotal_indices=subtotal_indices)

        self._update_regular_chart_cache(chart_root=chart_root, categories=categories, values=values)

        self.files[chart_part] = etree.tostring(chart_root, xml_declaration=True, encoding="UTF-8", standalone="yes")

        self._has_changes = True

    # Backward-compatible anywhere methods
    def edit_waterfall_data(self, chart_name: str, categories: list[str], values: list[float]) -> None:
        found = self.find_chart_anywhere(chart_name)
        self.edit_waterfall_data_on_slide(slide_number=found["slide_number"], chart_name=chart_name, categories=categories, values=values)

    def edit_embedded_workbook_for_chart(self, chart_name: str, categories: list[str], values: list[float], sheet_name: str | None = None, subtotal_indices: list[int] | None = None) -> None:
        found = self.find_chart_anywhere(chart_name)
        self.edit_embedded_workbook_for_chart_on_slide(slide_number=found["slide_number"], chart_name=chart_name, categories=categories, values=values, sheet_name=sheet_name, subtotal_indices=subtotal_indices)

    # PowerPoint COM refresh
    def _refresh_powerpoint_chart(self, output_pptx: str, chart_name: str) -> None:
        ppt_app = None
        presentation = None

        try:
            ppt_app = win32com.client.Dispatch("PowerPoint.Application")
            presentation = ppt_app.Presentations.Open(os.path.abspath(output_pptx))

            for slide in presentation.Slides:
                for shape in slide.Shapes:
                    if shape.Name == chart_name and shape.HasChart:
                        shape.Chart.Refresh()

            presentation.Save()

        except Exception as exc:
            print(f"Warning: Could not refresh chart in PowerPoint: {exc}")

        finally:
            try:
                if presentation is not None:
                    presentation.Close()
            except Exception:
                pass

            try:
                if ppt_app is not None:
                    ppt_app.Quit()
            except Exception:
                pass

    def refresh_chart(self, chart_name: str, output_pptx: str) -> None:
        self._refresh_powerpoint_chart(output_pptx, chart_name)

    # Debug/listing utilities
    def list_all_textboxes(self) -> None:
        slide_parts = [(int(match.group(1)), name) for name in self.files if (match := SLIDE_RE.match(name))]
        slide_parts.sort(key=lambda item: item[0])

        for slide_number, slide_part in slide_parts:
            slide_root = etree.fromstring(self.files[slide_part])

            for shape in slide_root.xpath(".//p:sp", namespaces=NS):
                cNvPr = shape.xpath("./p:nvSpPr/p:cNvPr", namespaces=NS)
                if not cNvPr:
                    continue

                name = cNvPr[0].get("name") or ""
                txBody = shape.xpath("./p:txBody", namespaces=NS)

                if txBody:
                    print(f"Slide {slide_number}: Textbox name={name!r}")

    def list_graphic_frames(self) -> None:
        slide_parts = [(int(match.group(1)), name) for name in self.files if (match := SLIDE_RE.match(name))]
        slide_parts.sort(key=lambda item: item[0])

        for slide_number, slide_part in slide_parts:
            slide_root = etree.fromstring(self.files[slide_part])
            frames = slide_root.xpath(".//p:graphicFrame", namespaces=NS)

            if not frames:
                continue

            print(f"\nSlide {slide_number}")

            for index, frame in enumerate(frames, start=1):
                cNvPr = frame.xpath("./p:nvGraphicFramePr/p:cNvPr", namespaces=NS)
                shape_name = cNvPr[0].get("name") if cNvPr else "<no name>"

                chart_refs = frame.xpath(".//c:chart", namespaces=NS)
                chartex_refs = frame.xpath(".//cx:chart", namespaces=NS)

                print(
                    f"  {index}. name={shape_name!r}, is_regular_chart={bool(chart_refs)}, is_chartex_chart={bool(chartex_refs)}"
                )

    def list_sections(self) -> None:
        presentation_root = etree.fromstring(self.files["ppt/presentation.xml"])
        sections = presentation_root.xpath(".//p14:section", namespaces=NS)

        if not sections:
            print("No sections found")
            return

        for index, section in enumerate(sections, start=1):
            name = section.get("name") or ""
            slide_ids = section.xpath("./p14:sldIdLst/p14:sldId", namespaces=NS)
            print(f"{index}. {name!r}: {len(slide_ids)} slides")

    def dump_chartex_debug(self, chart_name: str) -> None:
        found = self.find_chart_anywhere(chart_name)
        chart_part = found["chart_part"]
        chart_filename = posixpath.basename(chart_part)
        chart_dir = posixpath.dirname(chart_part)
        rels_part = f"{chart_dir}/_rels/{chart_filename}.rels"

        print("FOUND:", found)
        print()

        print("=== CHART PART ===")
        print(chart_part)
        print(self.files[chart_part].decode("utf-8", errors="ignore")[:8000])
        print()

        if rels_part in self.files:
            print("=== CHART RELS ===")
            print(rels_part)
            print(self.files[rels_part].decode("utf-8", errors="ignore")[:8000])
        else:
            print("No chart rels part found:", rels_part)

    # ============================================================
    # Table lookup and editing
    # ============================================================
    def find_table_on_slide(self, slide_number: int, table_name: str) -> dict:
        """Locate a table (graphicFrame containing an a:tbl) by its shape name on a slide.

        Returns a dict with slide_number and slide_part if found.
        """
        wanted = table_name.strip().casefold()

        slide_part = f"ppt/slides/slide{slide_number}.xml"
        if slide_part not in self.files:
            raise FileNotFoundError(f"Slide not found: {slide_part}")

        slide_root = etree.fromstring(self.files[slide_part])

        for frame in slide_root.xpath('.//p:graphicFrame', namespaces=NS):
            cNvPr = frame.xpath('./p:nvGraphicFramePr/p:cNvPr', namespaces=NS)
            if not cNvPr:
                continue

            actual_name = (cNvPr[0].get('name') or '').strip().casefold()
            if actual_name != wanted:
                continue

            # check if this graphicFrame contains a table
            tbl = frame.xpath('.//a:tbl', namespaces=NS)
            if not tbl:
                continue

            return {
                'slide_number': slide_number,
                'slide_part': slide_part,
            }

        raise ValueError(f"Table named '{table_name}' not found on slide {slide_number}")

    def find_table_anywhere(self, table_name: str) -> dict:
        slide_parts = [(int(match.group(1)), name) for name in self.files if (match := SLIDE_RE.match(name))]
        slide_parts.sort(key=lambda item: item[0])

        for slide_number, _slide_part in slide_parts:
            try:
                return self.find_table_on_slide(slide_number, table_name)
            except ValueError:
                continue

        raise ValueError(f"Table named '{table_name}' not found anywhere in presentation")

    def _set_table_cell_text(self, txBody_elem, text: str) -> None:
        """Helper: replace the text content of a table cell's txBody."""
        paragraphs = self._text_to_paragraph_runs(text)
        self._replace_textbox_runs(txBody_elem, paragraphs)

    def edit_table_cell_on_slide(self, slide_number: int, table_name: str, row: int, col: int, new_text: str) -> None:
        """Edit a single table cell by 0-based `row` and `col` indices on a slide."""
        info = self.find_table_on_slide(slide_number, table_name)
        slide_part = info['slide_part']

        slide_root = etree.fromstring(self.files[slide_part])

        # find the specific frame again and the table element
        for frame in slide_root.xpath('.//p:graphicFrame', namespaces=NS):
            cNvPr = frame.xpath('./p:nvGraphicFramePr/p:cNvPr', namespaces=NS)
            if not cNvPr:
                continue

            actual_name = (cNvPr[0].get('name') or '').strip().casefold()
            if actual_name != table_name.strip().casefold():
                continue

            tbls = frame.xpath('.//a:tbl', namespaces=NS)
            if not tbls:
                continue

            tbl = tbls[0]
            rows = tbl.xpath('./a:tr', namespaces=NS)
            if row < 0 or row >= len(rows):
                raise IndexError('row index out of range')

            cells = rows[row].xpath('./a:tc', namespaces=NS)
            if col < 0 or col >= len(cells):
                raise IndexError('col index out of range')

            cell = cells[col]
            txBody = cell.xpath('.//a:txBody', namespaces=NS)
            if not txBody:
                raise ValueError('Table cell has no text body')

            self._set_table_cell_text(txBody[0], new_text)

            # persist changes
            self.files[slide_part] = etree.tostring(slide_root, xml_declaration=True, encoding='UTF-8', standalone='yes')
            self._has_changes = True
            return

        raise ValueError(f"Table named '{table_name}' not found on slide {slide_number}")

    def edit_table_range_on_slide(self, slide_number: int, table_name: str, data: list[list[str]]) -> None:
        """Write a 2D list of strings into a table on a slide. Rows/cols must fit the table."""
        info = self.find_table_on_slide(slide_number, table_name)
        slide_part = info['slide_part']

        slide_root = etree.fromstring(self.files[slide_part])

        for frame in slide_root.xpath('.//p:graphicFrame', namespaces=NS):
            cNvPr = frame.xpath('./p:nvGraphicFramePr/p:cNvPr', namespaces=NS)
            if not cNvPr:
                continue

            actual_name = (cNvPr[0].get('name') or '').strip().casefold()
            if actual_name != table_name.strip().casefold():
                continue

            tbls = frame.xpath('.//a:tbl', namespaces=NS)
            if not tbls:
                continue

            tbl = tbls[0]
            rows = tbl.xpath('./a:tr', namespaces=NS)

            if len(data) > len(rows):
                raise ValueError('Provided data has more rows than the table')

            for r_idx, row_vals in enumerate(data):
                cells = rows[r_idx].xpath('./a:tc', namespaces=NS)
                if len(row_vals) > len(cells):
                    raise ValueError(f'Row {r_idx} has more columns than the table')

                for c_idx, val in enumerate(row_vals):
                    cell = cells[c_idx]
                    txBody = cell.xpath('.//a:txBody', namespaces=NS)
                    if not txBody:
                        raise ValueError('Table cell has no text body')
                    self._set_table_cell_text(txBody[0], val)

            # persist changes
            self.files[slide_part] = etree.tostring(slide_root, xml_declaration=True, encoding='UTF-8', standalone='yes')
            self._has_changes = True
            return

        raise ValueError(f"Table named '{table_name}' not found on slide {slide_number}")
